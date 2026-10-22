#Librerias
import customtkinter as ctk
from tkinter import filedialog, messagebox
import pandas as pd
import numpy as np
from decimal import Decimal, ROUND_HALF_UP
import openpyxl
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter
import datetime
import threading
import os

#Configuración visual de la ventana
ctk.set_appearance_mode("System")
ctk.set_default_color_theme("blue")

def redondear_excel(valor, decimales=0):
    """Replica exactamente la función REDONDEAR de Excel (ROUND_HALF_UP) exigida por la NOM."""
    if pd.isna(valor):
        return np.nan
    d = Decimal(str(valor))
    if decimales == 0:
        return int(d.quantize(Decimal('1'), rounding=ROUND_HALF_UP))
    else:
        formato = '0.' + '0' * decimales
        return float(d.quantize(Decimal(formato), rounding=ROUND_HALF_UP))

class ProcesadorApp(ctk.CTk):
    def __init__(self):
        super().__init__()
        
        self.title("Programax57")
        self.geometry("550x450")
        self.resizable(False, False)
        
        self.ruta_crudo = None
        self.ruta_ref = None
        
        #Título
        self.lbl_titulo = ctk.CTkLabel(self, text="programa57", font=ctk.CTkFont(size=20, weight="bold"))
        self.lbl_titulo.pack(pady=(20, 20))
        
        #Marco para el archivo Crudo
        self.frame_crudo = ctk.CTkFrame(self)
        self.frame_crudo.pack(pady=10, padx=20, fill="x")
        self.btn_crudo = ctk.CTkButton(self.frame_crudo, text="Cargar Excel Crudo", command=self.cargar_crudo)
        self.btn_crudo.pack(side="left", padx=10, pady=10)
        self.lbl_crudo_path = ctk.CTkLabel(self.frame_crudo, text="Ningún archivo seleccionado", text_color="gray")
        self.lbl_crudo_path.pack(side="left", padx=10)
        
        # Marco para el archivo de Referencia
        self.frame_ref = ctk.CTkFrame(self)
        self.frame_ref.pack(pady=10, padx=20, fill="x")
        self.btn_ref = ctk.CTkButton(self.frame_ref, text="Cargar Plantilla 2025", command=self.cargar_ref)
        self.btn_ref.pack(side="left", padx=10, pady=10)
        self.lbl_ref_path = ctk.CTkLabel(self.frame_ref, text="Ningún archivo seleccionado", text_color="gray")
        self.lbl_ref_path.pack(side="left", padx=10)
        
        # Botón de Procesar
        self.btn_procesar = ctk.CTkButton(self, text="Procesar y Generar Archivo", 
                                          command=self.iniciar_procesamiento, 
                                          font=ctk.CTkFont(size=15, weight="bold"),
                                          fg_color="green", hover_color="darkgreen")
        self.btn_procesar.pack(pady=(30, 10))
        
        # Etiqueta de estado
        self.lbl_estado = ctk.CTkLabel(self, text="Esperando archivos...", font=ctk.CTkFont(size=12))
        self.lbl_estado.pack(pady=10)

    def cargar_crudo(self):
        ruta = filedialog.askopenfilename(filetypes=[("Excel files", "*.xlsx *.xls")])
        if ruta:
            self.ruta_crudo = ruta
            nombre_archivo = os.path.basename(ruta)
            self.lbl_crudo_path.configure(text=nombre_archivo, text_color="white")
            
    def cargar_ref(self):
        ruta = filedialog.askopenfilename(filetypes=[("Excel files", "*.xlsx *.xls")])
        if ruta:
            self.ruta_ref = ruta
            nombre_archivo = os.path.basename(ruta)
            self.lbl_ref_path.configure(text=nombre_archivo, text_color="white")
            
    def iniciar_procesamiento(self):
        if not self.ruta_crudo or not self.ruta_ref:
            messagebox.showwarning("Faltan archivos", "Por favor selecciona ambos archivos antes de procesar.")
            return
            
        ruta_salida = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            initialfile="Resultado_Final.xlsx",
            title="Guardar archivo como",
            filetypes=[("Excel files", "*.xlsx")]
        )
        
        if not ruta_salida:
            return 
            
        self.btn_procesar.configure(state="disabled", text="Procesando...")
        self.lbl_estado.configure(text="Generando documento y tablas de perfiles...", text_color="yellow")
        
        hilo = threading.Thread(target=self.procesar_datos, args=(self.ruta_crudo, self.ruta_ref, ruta_salida))
        hilo.start()

    def procesar_datos(self, entrada, referencia, salida):
        try:
            # --- 1. LÓGICA DE PROCESAMIENTO CRUDO ---
            df_raw = pd.read_excel(entrada, header=2)
            df_datos = df_raw.iloc[1:].reset_index(drop=True)
            
            columnas_crudo = []
            parametros = []
            
            for i, col in enumerate(df_raw.columns):
                col_str = str(col)
                serie = df_datos[col].copy()

                if col_str.lower().startswith("status"):
                    serie.name = "Status"
                else:
                    serie.name = col

                columnas_crudo.append(serie)

                if str(col).lower().startswith("status") and i > 0:
                    param_name = df_raw.columns[i-1]
                    parametros.append(param_name)
                    
                    status_series = df_datos[col].astype(str).str.lower().str.strip()
                    valores_param = df_datos[param_name]
                    valores_numericos = pd.to_numeric(valores_param, errors='coerce')
                    condicion_valida = (status_series == 'ok') & (valores_numericos >= 0)
                    limpio_series = np.where(condicion_valida, valores_numericos, np.nan)
                    
                    col_limpio_nombre = f"{param_name}_Limpio"
                    columnas_crudo.append(pd.Series(limpio_series, name=col_limpio_nombre))
                    
            df_crudo = pd.concat(columnas_crudo, axis=1)
            
            df_ver = pd.DataFrame()
            df_ver['Date & Time'] = df_datos['Date & Time']
            dt = pd.to_datetime(df_datos['Date & Time'], format='%d/%m/%Y %H:%M', errors='coerce')
            df_ver['Date'] = dt.dt.strftime('%d/%m/%Y')
            df_ver['Hora'] = dt.dt.hour
            df_ver['Mes'] = dt.dt.month
            
            col_idx = 5 
            letras_valorvalidado = {} 
            
            for param in parametros:
                col_limpio = f"{param}_Limpio"
                col_val = f"{param}_validacion"
                col_validado = f"{param}_valorvalidado"
                
                df_ver[col_limpio] = df_crudo[col_limpio]
                letra_limpio = get_column_letter(col_idx)
                col_idx += 1
                
                df_ver[col_val] = ""
                letra_val = get_column_letter(col_idx)
                col_idx += 1
                
                letra_validado = get_column_letter(col_idx)
                letras_valorvalidado[param] = letra_validado
                col_idx += 1
                
                formulas_validado = []
                for row in range(2, len(df_datos) + 2):
                    formula = f'=IF({letra_limpio}{row}="","",IF({letra_val}{row}="IO","",{letra_limpio}{row}))'
                    formulas_validado.append(formula)
                df_ver[col_validado] = formulas_validado

            # --- 2. PREPARACIÓN Y CÁLCULOS NORMATIVOS (MOTOR NOM) ---
            nombres_mapeo_nom = {
                'PM10': 'PM10',
                'PM2_5': 'PM25',
                'PM2.5': 'PM25',
                'O3': 'O3',
                'SO2': 'SO2'
            }
            df_nom_base_dict = {'Fecha': df_ver['Date'], 'Hora': df_ver['Hora']}
            for p_orig in parametros:
                p_limpio_key = f"{p_orig}_Limpio"
                if p_limpio_key in df_crudo.columns:
                    nombre_destino = nombres_mapeo_nom.get(p_orig, p_orig)
                    df_nom_base_dict[nombre_destino] = pd.to_numeric(df_crudo[p_limpio_key], errors='coerce')

            df_base_nom = pd.DataFrame(df_nom_base_dict)
            for col_req in ['PM10', 'PM25', 'O3', 'SO2']:
                if col_req not in df_base_nom.columns:
                    df_base_nom[col_req] = np.nan

            try:
                fechas_dt = pd.to_datetime(df_base_nom['Fecha'], format='%d/%m/%Y')
                anio = int(fechas_dt.dt.year.mode()[0])
            except Exception:
                anio = 2025

            dias_lista = []
            for fecha_str, grupo in df_base_nom.groupby('Fecha', sort=False):
                try:
                    mes = int(pd.to_datetime(fecha_str, format='%d/%m/%Y').month)
                except Exception:
                    mes = 1
                trim = 1 if mes <= 3 else (2 if mes <= 6 else (3 if mes <= 9 else 4))

                dias_lista.append({
                    'Fecha': fecha_str, 'Mes': mes, 'Trimestre': trim,
                    'PM10': redondear_excel(grupo['PM10'].mean(), 0) if grupo['PM10'].count() >= 18 else np.nan,
                    'PM25': redondear_excel(grupo['PM25'].mean(), 0) if grupo['PM25'].count() >= 18 else np.nan,
                    'SO2_24h': redondear_excel(grupo['SO2'].mean(), 0) if grupo['SO2'].count() >= 18 else np.nan
                })
            df_diario = pd.DataFrame(dias_lista)

            dias_esperados = {1: 90, 2: 91, 3: 92, 4: 92}
            nombres_trim = {1: 'Ene-Mar', 2: 'Abr-Jun', 3: 'Jul-Sep', 4: 'Oct-Dic'}
            trim_lista = []
            for t in [1, 2, 3, 4]:
                sub = df_diario[df_diario['Trimestre'] == t] if not df_diario.empty else pd.DataFrame()
                esp = dias_esperados[t]
                if not sub.empty:
                    c10 = int((sub['PM10'] > 0).sum())
                    p10 = redondear_excel(sub['PM10'].mean(), 1) if c10 > 0 else np.nan
                    c25 = int((sub['PM25'] > 0).sum())
                    p25 = redondear_excel(sub['PM25'].mean(), 1) if c25 > 0 else np.nan
                else:
                    c10, p10, c25, p25 = 0, np.nan, 0, np.nan

                trim_lista.append({
                    'Conteo_esperado': esp, 'num': t, 'trimestre': nombres_trim[t],
                    'PM10_prom': p10, 'PM10_conteo': c10, 'PM10_prom75': p10 if c10 >= 0.75 * esp else np.nan,
                    'PM25_prom': p25, 'PM25_conteo': c25, 'PM25_prom75': p25 if c25 >= 0.75 * esp else np.nan
                })
            df_trim = pd.DataFrame(trim_lista)

            anual_pm10 = redondear_excel(df_trim['PM10_prom75'].mean(), 1) if df_trim['PM10_prom75'].count() >= 3 else np.nan
            anual_pm25 = redondear_excel(df_trim['PM25_prom75'].mean(), 1) if df_trim['PM25_prom75'].count() >= 3 else np.nan

            max_24h_pm10 = int(df_diario['PM10'].max()) if not df_diario.empty and not pd.isna(df_diario['PM10'].max()) else 0
            max_24h_pm25 = int(df_diario['PM25'].max()) if not df_diario.empty and not pd.isna(df_diario['PM25'].max()) else 0

            eval_24h_pm10 = "Incumple" if max_24h_pm10 > 75 else "Cumple"
            eval_anual_pm10 = "Incumple" if not pd.isna(anual_pm10) and anual_pm10 > 40 else "Cumple"
            resumen_pm10 = f"Incumple en {anio}" if "Incumple" in [eval_24h_pm10, eval_anual_pm10] else f"Cumple en {anio}"

            eval_24h_pm25 = "Incumple" if max_24h_pm25 > 45 else "Cumple"
            eval_anual_pm25 = "Incumple" if not pd.isna(anual_pm25) and anual_pm25 > 12 else "Cumple"
            resumen_pm25 = f"Incumple en {anio}" if "Incumple" in [eval_24h_pm25, eval_anual_pm25] else f"Cumple en {anio}"

            o3_mov8, so2_mov8 = [], []
            vals_o3, vals_so2 = df_base_nom['O3'].values, df_base_nom['SO2'].values
            for i in range(len(vals_o3)):
                if i < 5:
                    o3_mov8.append(np.nan)
                    so2_mov8.append(np.nan)
                else:
                    v_o = vals_o3[max(0, i - 7):i + 1]
                    o3_mov8.append(redondear_excel(np.nanmean(v_o), 0) if np.sum(~pd.isna(v_o)) >= 6 else np.nan)
                    v_s = vals_so2[max(0, i - 7):i + 1]
                    so2_mov8.append(redondear_excel(np.nanmean(v_s), 0) if np.sum(~pd.isna(v_s)) >= 6 else np.nan)

            o3_1h_raw = redondear_excel(df_base_nom['O3'].max(), 0) if df_base_nom['O3'].count() > 0 else 0
            o3_max_1h = o3_1h_raw if (o3_1h_raw > 95 or df_base_nom['O3'].count() >= 0.75 * 8760) else "DI"
            eval_o3_1h = "DI" if o3_max_1h == "DI" else ("Incumple" if o3_max_1h > 95 else "Cumple")

            s_o3_mov8 = pd.Series(o3_mov8)
            o3_8h_raw = redondear_excel(s_o3_mov8.max(), 0) if s_o3_mov8.count() > 0 else 0
            o3_max_8h = o3_8h_raw if (o3_8h_raw > 70 or s_o3_mov8.count() >= 0.75 * 8760) else "DI"
            eval_o3_8h = "DI" if o3_max_8h == "DI" else ("Incumple" if o3_max_8h > 70 else "Cumple")
            resumen_o3 = f"O3 incumple en el año {anio}" if "Incumple" in [eval_o3_1h, eval_o3_8h] else f"O3 cumple en el año {anio}"

            so2_anual_val = redondear_excel(df_base_nom['SO2'].mean(), 0) if df_base_nom['SO2'].count() >= 0.75 * 8760 else "DI"
            eval_so2_anual = "DI" if so2_anual_val == "DI" else ("Incumple" if so2_anual_val > 25 else "Cumple")

            s_so2_mov8 = pd.Series(so2_mov8)
            so2_8h_raw = redondear_excel(s_so2_mov8.max(), 0) if s_so2_mov8.count() > 0 else 0
            so2_max_8h = so2_8h_raw if (so2_8h_raw > 200 or s_so2_mov8.count() >= 0.75 * 8760) else "DI"
            eval_so2_8h = "DI" if so2_max_8h == "DI" else ("Incumple" if so2_max_8h > 200 else "Cumple")

            so2_24h_raw = int(df_diario['SO2_24h'].max()) if not df_diario.empty and not pd.isna(df_diario['SO2_24h'].max()) else 0
            so2_max_24h = so2_24h_raw if (so2_24h_raw > 110 or df_diario['SO2_24h'].count() >= 0.75 * 365) else "DI"
            eval_so2_24h = "DI" if so2_max_24h == "DI" else ("Incumple" if so2_max_24h > 110 else "Cumple")

            # --- 3. ESCRITURA TOTALMENTE DINÁMICA DE EXCEL ---
            self.lbl_estado.configure(text="Diseñando tablas y perfiles desde cero...", text_color="orange")
            
            with pd.ExcelWriter(salida, engine='openpyxl') as writer:
                df_crudo.to_excel(writer, sheet_name='Crudo', index=False)
                try:
                    df_flags = pd.read_excel(referencia, sheet_name='FLAGS')
                    df_flags.to_excel(writer, sheet_name='FLAGS', index=False)
                except Exception:
                    pass
                df_ver.to_excel(writer, sheet_name='Verificación', index=False)
                pd.DataFrame().to_excel(writer, sheet_name='VSAFE', index=False)
                
                # --- MODIFICACIONES EN LA HOJA VERIFICACIÓN ---
                workbook = writer.book
                ws_ver = workbook['Verificación']
                
                negrita = Font(bold=True)

                # Creación del Perfil Horario
                col_horario = 62 # Columna BJ
                ws_ver.cell(row=1, column=col_horario, value="Perfil horario").font = negrita
                ws_ver.cell(row=2, column=col_horario, value="Hora").font = negrita
                
                df_horario = df_crudo[[f"{p}_Limpio" for p in parametros]].copy()
                df_horario['Hora'] = df_ver['Hora']
                df_horario_mean = df_horario.groupby('Hora').mean(numeric_only=True).reset_index()
                
                for r_idx, hora in enumerate(df_horario_mean['Hora']):
                    c = ws_ver.cell(row=r_idx+3, column=col_horario, value=datetime.time(hour=int(hora)))
                    c.number_format = 'hh:mm'
                    
                for i, p in enumerate(parametros):
                    col_p = col_horario + 1 + i
                    ws_ver.cell(row=2, column=col_p, value=p).font = negrita
                    for r_idx, val in enumerate(df_horario_mean[f"{p}_Limpio"]):
                        if not pd.isna(val):
                            ws_ver.cell(row=r_idx+3, column=col_p, value=val)

                # Creación del Perfil Mensual
                col_mensual = col_horario + len(parametros) + 2
                ws_ver.cell(row=1, column=col_mensual, value="Perfil mensual").font = negrita
                
                encabezados_m = ["Mes", "N° MES", "Datos esp. Por mes"]
                for i, texto in enumerate(encabezados_m):
                    ws_ver.cell(row=2, column=col_mensual+i, value=texto).font = negrita
                    
                df_mensual_base = df_crudo[[f"{p}_Limpio" for p in parametros]].copy()
                df_mensual_base['Mes'] = df_ver['Mes']
                df_mensual_base['Date'] = dt
                meses_presentes = df_mensual_base['Mes'].dropna().unique()
                meses_presentes.sort()
                nombres_meses = ["Ene", "Feb", "Mar", "Abr", "May", "Jun", "Jul", "Ago", "Sep", "Oct", "Nov", "Dic"]
                
                for r_idx, mes in enumerate(meses_presentes):
                    mes = int(mes)
                    row_idx = r_idx + 3
                    ws_ver.cell(row=row_idx, column=col_mensual, value=nombres_meses[mes-1] if 1<=mes<=12 else str(mes))
                    ws_ver.cell(row=row_idx, column=col_mensual+1, value=mes)
                    
                    try:
                        anio_mes = df_mensual_base[df_mensual_base['Mes'] == mes]['Date'].dt.year.mode()[0]
                        dias_mes = pd.Period(f"{int(anio_mes)}-{mes:02d}").days_in_month
                    except Exception:
                        dias_mes = 30
                    datos_esp = dias_mes * 24
                    ws_ver.cell(row=row_idx, column=col_mensual+2, value=datos_esp)
                    
                    c_p = col_mensual + 3
                    for p in parametros:
                        ws_ver.cell(row=2, column=c_p, value=p).font = negrita
                        ws_ver.cell(row=2, column=c_p+1, value="Datos válidos por mes").font = negrita
                        ws_ver.cell(row=2, column=c_p+2, value="% Datos Validos").font = negrita
                        
                        data_mes = df_mensual_base[df_mensual_base['Mes'] == mes][f"{p}_Limpio"]
                        prom = data_mes.mean()
                        count = data_mes.count()
                        pct = count / datos_esp if datos_esp > 0 else 0
                        
                        if not pd.isna(prom): ws_ver.cell(row=row_idx, column=c_p, value=prom)
                        ws_ver.cell(row=row_idx, column=c_p+1, value=count)
                        c_pct = ws_ver.cell(row=row_idx, column=c_p+2, value=pct)
                        c_pct.number_format = '0.00%'
                        c_p += 3

                # --- 4. CREAR HOJA NOM ---
                # --- 4. CREAR HOJA NOM (100% DINÁMICA CON FÓRMULAS DE EXCEL) ---
                ws = workbook.create_sheet(title='NOM')
                
                # Mapeo de nombres posibles en el archivo crudo hacia los 4 parámetros de la NOM
                variantes_param = {
                    'PM10': ['PM10'],
                    'PM25': ['PM2_5', 'PM2.5', 'PM25'],
                    'O3': ['O3'],
                    'SO2': ['SO2']
                }
                letras_nom_ver = {}
                for nom_key, lista_posibles in variantes_param.items():
                    for p_posible in lista_posibles:
                        if p_posible in letras_valorvalidado:
                            letras_nom_ver[nom_key] = letras_valorvalidado[p_posible]
                            break

                try:
                    anio = int(dt.dt.year.mode()[0])
                except Exception:
                    anio = 2025

                num_filas_horarias = len(df_ver)
                max_row_h = num_filas_horarias + 1

                # TABLA 1: Datos horarios validados (Vinculados a 'Verificación')
                encabezados_af = ['Fecha', 'Hora', 'PM10', 'PM25', 'O3', 'SO2']
                for col_idx, texto in enumerate(encabezados_af, start=1):
                    ws.cell(row=1, column=col_idx, value=texto)

                for r_idx in range(num_filas_horarias):
                    fila_excel = r_idx + 2
                    ws.cell(row=fila_excel, column=1, value=df_ver['Date'].iloc[r_idx])
                    hora_val = df_ver['Hora'].iloc[r_idx]
                    ws.cell(row=fila_excel, column=2, value=int(hora_val) if not pd.isna(hora_val) else 0)
                    
                    for c_i, nombre_col in enumerate(['PM10', 'PM25', 'O3', 'SO2'], start=3):
                        if nombre_col in letras_nom_ver:
                            letra_ver = letras_nom_ver[nombre_col]
                            # Si en Verificación está vacío (por "IO" o falta de dato), deja "", si no, pasa el número
                            ws.cell(row=fila_excel, column=c_i, value=f'=IF(\'Verificación\'!{letra_ver}{fila_excel}="","",\'Verificación\'!{letra_ver}{fila_excel})')

                # TABLA 3: Promedio de 24 hr filtrado (Calculado dinámicamente desde Tabla 1)
                ws['M1'] = "Promedio de 24 hr"
                ws['M1'].font = negrita
                # Se incluye SO2_24h en la columna S (19) para que el cálculo de SO2 24h también sea dinámico
                for offset, texto in enumerate(['Fecha', 'Mes', 'Trimestre', 'PM10', 'PM25', 'SO2_24h']):
                    ws.cell(row=1, column=14 + offset, value=texto)

                fechas_unicas = df_ver['Date'].dropna().unique()
                num_dias = len(fechas_unicas)
                max_row_d = num_dias + 1

                for r_idx, fecha_str in enumerate(fechas_unicas):
                    fila_d = r_idx + 2
                    try:
                        mes = int(pd.to_datetime(fecha_str, format='%d/%m/%Y').month)
                    except Exception:
                        mes = 1
                    trim = 1 if mes <= 3 else (2 if mes <= 6 else (3 if mes <= 9 else 4))

                    ws.cell(row=fila_d, column=14, value=fecha_str) # Col N: Fecha
                    ws.cell(row=fila_d, column=15, value=mes)       # Col O: Mes
                    ws.cell(row=fila_d, column=16, value=trim)      # Col P: Trimestre

                    # Col Q (17): Promedio 24h PM10 (Col C) si hay >= 18 datos válidos en ese día
                    ws.cell(row=fila_d, column=17, value=f'=IF(COUNTIFS($A$2:$A${max_row_h}, N{fila_d}, $C$2:$C${max_row_h}, ">=0")>=18, ROUND(AVERAGEIFS($C$2:$C${max_row_h}, $A$2:$A${max_row_h}, N{fila_d}, $C$2:$C${max_row_h}, ">=0"), 0), "")')
                    # Col R (18): Promedio 24h PM25 (Col D) si hay >= 18 datos válidos en ese día
                    ws.cell(row=fila_d, column=18, value=f'=IF(COUNTIFS($A$2:$A${max_row_h}, N{fila_d}, $D$2:$D${max_row_h}, ">=0")>=18, ROUND(AVERAGEIFS($D$2:$D${max_row_h}, $A$2:$A${max_row_h}, N{fila_d}, $D$2:$D${max_row_h}, ">=0"), 0), "")')
                    # Col S (19): Promedio 24h SO2 (Col F) si hay >= 18 datos válidos en ese día
                    ws.cell(row=fila_d, column=19, value=f'=IF(COUNTIFS($A$2:$A${max_row_h}, N{fila_d}, $F$2:$F${max_row_h}, ">=0")>=18, ROUND(AVERAGEIFS($F$2:$F${max_row_h}, $A$2:$A${max_row_h}, N{fila_d}, $F$2:$F${max_row_h}, ">=0"), 0), "")')

                # TABLA 4: Promedios trimestrales y anual
                ws['T1'], ws['T2'], ws['V2'] = "Promedios trimestrales", "Conteo esperado", "trimestre"
                ws['T1'].font = negrita
                ws['W1'], ws['W2'], ws['X2'], ws['Y2'] = "PM10", "promedio", "conteo", "Promedio con 75%"
                ws['Z1'], ws['Z2'], ws['AA2'], ws['AB2'] = "PM2.5", "promedio", "conteo", "Promedio con 75%"

                dias_esperados = {1: 90, 2: 91, 3: 92, 4: 92}
                nombres_trim = {1: 'Ene-Mar', 2: 'Abr-Jun', 3: 'Jul-Sep', 4: 'Oct-Dic'}

                for idx, t in enumerate([1, 2, 3, 4]):
                    r = idx + 3
                    ws.cell(row=r, column=20, value=dias_esperados[t]) # Col T: Conteo esperado
                    ws.cell(row=r, column=21, value=t)                 # Col U: Num trimestre
                    ws.cell(row=r, column=22, value=nombres_trim[t])   # Col V: Nombre trimestre

                    # PM10 Trimestral (Col W: promedio, Col X: conteo, Col Y: Promedio con 75%)
                    ws.cell(row=r, column=23, value=f'=IF(X{r}>0, ROUND(AVERAGEIFS($Q$2:$Q${max_row_d}, $P$2:$P${max_row_d}, U{r}, $Q$2:$Q${max_row_d}, ">=0"), 1), "")')
                    ws.cell(row=r, column=24, value=f'=COUNTIFS($P$2:$P${max_row_d}, U{r}, $Q$2:$Q${max_row_d}, ">0")')
                    ws.cell(row=r, column=25, value=f'=IF(X{r}>=0.75*T{r}, W{r}, "")')

                    # PM2.5 Trimestral (Col Z: promedio, Col AA: conteo, Col AB: Promedio con 75%)
                    ws.cell(row=r, column=26, value=f'=IF(AA{r}>0, ROUND(AVERAGEIFS($R$2:$R${max_row_d}, $P$2:$P${max_row_d}, U{r}, $R$2:$R${max_row_d}, ">=0"), 1), "")')
                    ws.cell(row=r, column=27, value=f'=COUNTIFS($P$2:$P${max_row_d}, U{r}, $R$2:$R${max_row_d}, ">0")')
                    ws.cell(row=r, column=28, value=f'=IF(AA{r}>=0.75*T{r}, Z{r}, "")')

                ws['V8'] = "Promedio anual"
                ws['V8'].font = negrita
                ws['Y8'] = '=IF(COUNT(Y3:Y6)>=3, ROUND(AVERAGE(Y3:Y6), 1), "")'
                ws['Y8'].font = negrita
                ws['AB8'] = '=IF(COUNT(AB3:AB6)>=3, ROUND(AVERAGE(AB3:AB6), 1), "")'
                ws['AB8'].font = negrita

                # TABLA 2: Evaluación PM10 y PM2.5
                ws['H1'], ws['H2'], ws['H3'], ws['H6'] = "Evaluación PM10 y PM2.5", "NOM-025-SSA1-2014", "PM10", "PM2.5"
                for celda in ['H1', 'H2', 'H3', 'H6']: ws[celda].font = negrita
                ws['H4'], ws['I4'], ws['H5'], ws['I5'] = "24 horas", 75, "anual", 40
                ws['H7'], ws['I7'], ws['H8'], ws['I8'] = "24 horas", 45, "anual", 12

                ws['J3'], ws['K3'] = f"Año {anio} (promedio)", "Evaluación"
                ws['J3'].font, ws['K3'].font = negrita, negrita
                
                # Fórmulas de evaluación PM10
                ws['J4'] = f'=MAX(Q2:Q{max_row_d})'
                ws['K4'] = '=IF(J4>I4, "Incumple", "Cumple")'
                ws['J5'] = '=Y8'
                ws['K5'] = '=IF(J5="", "Cumple", IF(J5>I5, "Incumple", "Cumple"))'
                
                # Fórmulas de evaluación PM2.5
                ws['J7'] = f'=MAX(R2:R{max_row_d})'
                ws['K7'] = '=IF(J7>I7, "Incumple", "Cumple")'
                ws['J8'] = '=AB8'
                ws['K8'] = '=IF(J8="", "Cumple", IF(J8>I8, "Incumple", "Cumple"))'

                ws['H10'], ws['H11'] = "PM10", f'=IF(OR(K4="Incumple", K5="Incumple"), "Incumple en {anio}", "Cumple en {anio}")'
                ws['H13'], ws['H14'] = "PM2.5", f'=IF(OR(K7="Incumple", K8="Incumple"), "Incumple en {anio}", "Cumple en {anio}")'
                ws['H10'].font, ws['H13'].font = negrita, negrita

                # TABLA 5 y 6: Promedios móviles de 8h y Evaluación O3 y SO2
                ws['AJ1'], ws['AK1'], ws['AL1'], ws['AM1'] = "Promedio móvil de 8h", "Fecha", "Hora", "O3"
                ws['AJ1'].font = negrita
                ws['AT1'], ws['AU1'], ws['AV1'], ws['AW1'] = "Promedio móvil de 8 hrs", "Fecha", "Hora", "SO2"
                ws['AT1'].font = negrita

                for r_idx in range(num_filas_horarias):
                    fila_excel = r_idx + 2
                    # Fechas y horas vinculadas a Tabla 1
                    ws.cell(row=fila_excel, column=37, value=f'=A{fila_excel}') # AK
                    ws.cell(row=fila_excel, column=38, value=f'=B{fila_excel}') # AL
                    ws.cell(row=fila_excel, column=47, value=f'=A{fila_excel}') # AU
                    ws.cell(row=fila_excel, column=48, value=f'=B{fila_excel}') # AV

                    if r_idx >= 5:
                        inicio_ventana = max(2, fila_excel - 7)
                        # Col AM (39): Promedio móvil 8h O3 (Col E) si hay >= 6 datos en la ventana
                        ws.cell(row=fila_excel, column=39, value=f'=IF(COUNT(E{inicio_ventana}:E{fila_excel})>=6, ROUND(AVERAGE(E{inicio_ventana}:E{fila_excel}), 0), "")')
                        # Col AW (49): Promedio móvil 8h SO2 (Col F) si hay >= 6 datos en la ventana
                        ws.cell(row=fila_excel, column=49, value=f'=IF(COUNT(F{inicio_ventana}:F{fila_excel})>=6, ROUND(AVERAGE(F{inicio_ventana}:F{fila_excel}), 0), "")')

                # Cuadro Evaluación O3
                ws['AF1'], ws['AG1'], ws['AH1'] = "Evaluación O3", "Máx", "Evaluación"
                ws['AF1'].font = negrita
                ws['AF2'], ws['AF3'] = "1hr", 95
                ws['AG2'] = f'=IF(OR(IF(COUNT(E2:E{max_row_h})>0, ROUND(MAX(E2:E{max_row_h}), 0), 0)>AF3, COUNT(E2:E{max_row_h})>=6570), IF(COUNT(E2:E{max_row_h})>0, ROUND(MAX(E2:E{max_row_h}), 0), 0), "DI")'
                ws['AH2'] = '=IF(AG2="DI", "DI", IF(AG2>AF3, "Incumple", "Cumple"))'

                ws['AF4'], ws['AF5'] = "móvil de 8 hr", 70
                ws['AG5'] = f'=IF(OR(MAX(AM2:AM{max_row_h})>AF5, COUNT(AM2:AM{max_row_h})>=6570), MAX(AM2:AM{max_row_h}), "DI")'
                ws['AH5'] = '=IF(AG5="DI", "DI", IF(AG5>AF5, "Incumple", "Cumple"))'

                ws['AF7'], ws['AF8'] = "Evaluación", f'=IF(OR(AH2="Incumple", AH5="Incumple"), "O3 incumple en el año {anio}", "O3 cumple en el año {anio}")'
                ws['AF7'].font = negrita

                # Cuadro Evaluación SO2
                ws['AP1'], ws['AQ1'], ws['AR1'] = "Evaluación SO2", "Promedio", "Evaluación"
                ws['AP1'].font = negrita
                ws['AP2'], ws['AP3'] = "8 hr", 200
                ws['AQ3'] = f'=IF(OR(MAX(AW2:AW{max_row_h})>AP3, COUNT(AW2:AW{max_row_h})>=6570), MAX(AW2:AW{max_row_h}), "DI")'
                ws['AR3'] = '=IF(AQ3="DI", "DI", IF(AQ3>AP3, "Incumple", "Cumple"))'

                ws['AP4'], ws['AP5'] = "24 hrs", 110
                ws['AQ5'] = f'=IF(OR(MAX(S2:S{max_row_d})>AP5, COUNT(S2:S{max_row_d})>=273.75), MAX(S2:S{max_row_d}), "DI")'
                ws['AR5'] = '=IF(AQ5="DI", "DI", IF(AQ5>AP5, "Incumple", "Cumple"))'

                ws['AP6'], ws['AP7'] = "anual", 25
                ws['AQ7'] = f'=IF(COUNT(F2:F{max_row_h})>=6570, ROUND(AVERAGE(F2:F{max_row_h}), 0), "DI")'
                ws['AR7'] = '=IF(AQ7="DI", "DI", IF(AQ7>AP7, "Incumple", "Cumple"))'

                # --- 5. FORMATO FINAL PARA TODAS LAS HOJAS ---
                # Aplica negrita y centrado a todas las celdas de la primera fila en cualquier hoja generada
                for sheet in workbook.worksheets:
                    for cell in sheet[1]:
                        cell.font = Font(bold=True)
                        cell.alignment = Alignment(horizontal='center', vertical='center')

            self.lbl_estado.configure(text=f"¡Completado! Guardado en:\n{os.path.basename(salida)}", text_color="green")
            messagebox.showinfo("Éxito", "El archivo se ha procesado con las tablas de perfiles sin color.")
            
        except Exception as e:
            self.lbl_estado.configure(text="Ocurrió un error.", text_color="red")
            messagebox.showerror("Error", f"Detalle del error:\n{str(e)}")
            
        finally:
            self.btn_procesar.configure(state="normal", text="Procesar y Generar Archivo")

if __name__ == "__main__":
    app = ProcesadorApp()
    app.mainloop()